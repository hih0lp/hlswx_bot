from __future__ import annotations

import json
import re
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path

import joblib
import openpyxl
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import classification_report
from sklearn.model_selection import train_test_split
from sklearn.pipeline import Pipeline
from sqlalchemy import select

from app.db.session import SessionLocal
from app.ml.categories import category_from_price
from app.ml.preprocess import normalize_text
from app.models.entities import TrainingSample

REJECT_MARKERS = ("отказ",)


def _norm_price(value) -> float | str | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().lower()
    if any(m in s for m in REJECT_MARKERS):
        return "REJECT"
    m = re.search(r"(\d+)", s)
    return float(m.group(1)) if m else None


def load_rows_from_excel(path: Path) -> list[tuple[str, str]]:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows: list[tuple[str, str]] = []
    for row in ws.iter_rows(values_only=True):
        if not row or not row[0]:
            continue
        text = str(row[0]).strip()
        if text.upper().startswith("ЕЖЕДНЕВНЫЕ"):
            continue
        price = _norm_price(row[2] if len(row) > 2 else None)
        if price == "REJECT":
            rows.append((text, "REJECT"))
        elif isinstance(price, float):
            rows.append((text, category_from_price(price)))
    wb.close()
    return rows


async def load_rows_from_db() -> list[tuple[str, str]]:
    async with SessionLocal() as session:
        samples = (
            await session.scalars(select(TrainingSample).order_by(TrainingSample.id))
        ).all()
    return [(s.text, s.target_category) for s in samples]


def build_training_corpus(excel_path: Path, db_rows: list[tuple[str, str]]) -> list[tuple[str, str]]:
    merged: dict[str, str] = {}
    for text, cat in load_rows_from_excel(excel_path):
        merged[text] = cat
    for text, cat in db_rows:
        merged[text] = cat
    return list(merged.items())


def train_classifier(rows: list[tuple[str, str]]) -> tuple[Pipeline, dict]:
    if len(rows) < 20:
        raise RuntimeError(f"Not enough training rows: {len(rows)}")

    texts = [normalize_text(t) for t, _ in rows]
    labels = [cat for _, cat in rows]
    label_set = sorted(set(labels))

    counts = Counter(labels)
    stratify = labels if len(counts) > 1 and min(counts.values()) >= 2 else None
    X_train, X_test, y_train, y_test = train_test_split(
        texts,
        labels,
        test_size=0.15,
        random_state=42,
        stratify=stratify,
    )

    base_clf = LogisticRegression(max_iter=4000, class_weight="balanced")
    pipeline = Pipeline(
        [
            ("tfidf", TfidfVectorizer(ngram_range=(1, 2), max_features=15000, min_df=1)),
            ("clf", base_clf),
        ]
    )
    pipeline.fit(X_train, y_train)

    y_pred = pipeline.predict(X_test)
    report = classification_report(y_test, y_pred, zero_division=0, output_dict=True)
    metrics = {
        "trained_at": datetime.now(UTC).isoformat(),
        "samples": len(rows),
        "labels": label_set,
        "accuracy": report.get("accuracy", 0),
        "macro_f1": report.get("macro avg", {}).get("f1-score", 0),
    }
    return pipeline, metrics


def save_model(pipeline: Pipeline, model_path: Path, metrics: dict) -> None:
    model_path.parent.mkdir(parents=True, exist_ok=True)
    joblib.dump(pipeline, model_path)
    meta_path = model_path.with_suffix(".json")
    meta_path.write_text(json.dumps(metrics, ensure_ascii=False, indent=2), encoding="utf-8")


async def train_and_save(excel_path: Path, model_path: Path) -> dict:
    db_rows = await load_rows_from_db()
    rows = build_training_corpus(excel_path, db_rows)
    pipeline, metrics = train_classifier(rows)
    save_model(pipeline, model_path, metrics)
    return metrics
