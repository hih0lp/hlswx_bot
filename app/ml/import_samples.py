from __future__ import annotations

import csv
import json
import zipfile
from io import BytesIO, StringIO
from pathlib import Path

import openpyxl

from app.ml.categories import CATEGORIES, category_from_price
from app.ml.train import _norm_price, load_rows_from_excel

TRAINING_CATEGORY_CODES = frozenset(
    code for code, cat in CATEGORIES.items() if code not in ("INCOMPLETE",)
)
HEADER_MARKERS = ("текст", "text", "объявление", "категория", "category", "цена", "price", "label")

SUPPORTED_EXTENSIONS = (
    ".xlsx",
    ".xlsm",
    ".csv",
    ".tsv",
    ".txt",
    ".json",
    ".jsonl",
    ".ndjson",
    ".zip",
    ".jpg",
    ".jpeg",
    ".png",
    ".webp",
    ".bmp",
    ".gif",
)


def supported_formats_hint() -> str:
    return (
        "<b>Файлы:</b> .xlsx · .csv · .tsv · .txt · .json · .jsonl · .zip\n"
        "<b>Картинки:</b> .jpg/.png/.webp или фото в чат (OCR)\n\n"
        "<b>Excel:</b> A — текст, B или C — категория/цена\n"
        "<b>CSV/TSV/TXT:</b> <code>VACANCY|текст</code> или <code>текст;1200</code>\n"
        "<b>JSON:</b> <code>[{\"text\": \"...\", \"category\": \"VACANCY\"}]</code>\n"
        "<b>ZIP:</b> manifest (.json/.csv/.xlsx) + опционально папка images/\n"
        "<b>Фото:</b> подпись <code>VACANCY</code> или <code>VACANCY\\nтекст</code>"
    )


def resolve_category_token(token: str) -> str | None:
    raw = (token or "").strip()
    if not raw:
        return None
    upper = raw.upper()
    if upper in TRAINING_CATEGORY_CODES:
        return upper
    lower = raw.lower()
    if any(marker in lower for marker in ("отказ", "reject")):
        return "REJECT"
    price = _norm_price(raw)
    if price == "REJECT":
        return "REJECT"
    if isinstance(price, float):
        return category_from_price(price)
    return None


def _append_row(rows: list[tuple[str, str]], text: str, category: str) -> None:
    clean = (text or "").strip()
    cat = (category or "").strip()
    if len(clean) < 15 or not cat:
        return
    rows.append((clean, cat))


def _pair_from_tokens(left: str, right: str) -> tuple[str, str] | None:
    left_cat = resolve_category_token(left)
    right_cat = resolve_category_token(right)
    if left_cat and not right_cat and len(right.strip()) >= 15:
        return right.strip(), left_cat
    if right_cat and not left_cat and len(left.strip()) >= 15:
        return left.strip(), right_cat
    if left_cat and len(right.strip()) >= 15:
        return right.strip(), left_cat
    if right_cat and len(left.strip()) >= 15:
        return left.strip(), right_cat
    return None


def _row_from_mapping(item: dict) -> tuple[str, str] | None:
    if not isinstance(item, dict):
        return None
    text = (
        item.get("text")
        or item.get("advert")
        or item.get("message")
        or item.get("content")
        or item.get("body")
    )
    if not text:
        return None
    text = str(text).strip()
    for key in ("category", "category_code", "label", "class", "target"):
        if key in item and item[key]:
            cat = resolve_category_token(str(item[key]))
            if cat:
                return text, cat
    if "price" in item:
        cat = resolve_category_token(str(item["price"]))
        if cat:
            return text, cat
    return None


def load_rows_from_excel_bytes(data: bytes) -> list[tuple[str, str]]:
    wb = openpyxl.load_workbook(BytesIO(data), read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows: list[tuple[str, str]] = []
    for row in ws.iter_rows(values_only=True):
        if not row or not row[0]:
            continue
        text = str(row[0]).strip()
        if text.upper().startswith("ЕЖЕДНЕВНЫЕ"):
            continue
        if text.lower() in HEADER_MARKERS:
            continue
        price = _norm_price(row[2] if len(row) > 2 else None)
        if price == "REJECT":
            _append_row(rows, text, "REJECT")
            continue
        if isinstance(price, float):
            _append_row(rows, text, category_from_price(price))
            continue
        if len(row) > 1:
            category = resolve_category_token(str(row[1]))
            if category:
                _append_row(rows, text, category)
                continue
        if len(row) > 2:
            category = resolve_category_token(str(row[2]))
            if category:
                _append_row(rows, text, category)
    wb.close()
    return rows


def load_rows_from_csv_bytes(data: bytes, delimiter: str = ",") -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    content = None
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            content = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if content is None:
        raise ValueError("Не удалось прочитать CSV (кодировка)")

    reader = csv.reader(StringIO(content), delimiter=delimiter)
    for raw_row in reader:
        if not raw_row:
            continue
        cells = [c.strip() for c in raw_row if c and str(c).strip()]
        if not cells:
            continue
        if len(cells) == 1:
            line = cells[0]
            if line.lower() in HEADER_MARKERS or line.startswith("#"):
                continue
            for sep in ("|", ";", "\t"):
                if sep in line:
                    left, right = line.split(sep, 1)
                    pair = _pair_from_tokens(left, right)
                    if pair:
                        rows.append(pair)
                    break
            continue
        if cells[0].lower() in HEADER_MARKERS:
            continue
        if len(cells) >= 2:
            pair = _pair_from_tokens(cells[0], cells[1])
            if pair:
                rows.append(pair)
                continue
            if len(cells) >= 3:
                text = cells[0]
                price = _norm_price(cells[2])
                if isinstance(price, float):
                    _append_row(rows, text, category_from_price(price))
                elif price == "REJECT":
                    _append_row(rows, text, "REJECT")
    return rows


def load_rows_from_delimited_text(content: str) -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    for line in content.splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        if line.lower() in HEADER_MARKERS:
            continue
        matched = False
        for sep in ("|", ";", "\t"):
            if sep not in line:
                continue
            left, right = line.split(sep, 1)
            pair = _pair_from_tokens(left, right)
            if pair:
                rows.append(pair)
            matched = True
            break
        if not matched and "," in line and "|" not in line:
            pair = None
            try:
                cells = next(csv.reader(StringIO(line)))
                if len(cells) >= 2:
                    pair = _pair_from_tokens(cells[0], cells[1])
            except csv.Error:
                pair = None
            if pair:
                rows.append(pair)
    return rows


def load_rows_from_json_bytes(data: bytes) -> list[tuple[str, str]]:
    raw = data.decode("utf-8-sig")
    rows: list[tuple[str, str]] = []
    try:
        payload = json.loads(raw)
    except json.JSONDecodeError:
        return load_rows_from_jsonl_bytes(data)

    items: list = []
    if isinstance(payload, list):
        items = payload
    elif isinstance(payload, dict):
        for key in ("samples", "data", "rows", "items", "training"):
            if isinstance(payload.get(key), list):
                items = payload[key]
                break
        if not items and ("text" in payload or "category" in payload):
            items = [payload]
    for item in items:
        pair = _row_from_mapping(item) if isinstance(item, dict) else None
        if pair:
            rows.append(pair)
    return rows


def load_rows_from_jsonl_bytes(data: bytes) -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    content = data.decode("utf-8-sig")
    for line in content.splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        try:
            item = json.loads(line)
        except json.JSONDecodeError:
            continue
        pair = _row_from_mapping(item) if isinstance(item, dict) else None
        if pair:
            rows.append(pair)
    return rows


def load_rows_from_image_bytes(data: bytes, filename: str) -> list[tuple[str, str]]:
    from app.ml.image_ocr import (
        category_from_image_stem,
        is_image_filename,
        ocr_image_bytes,
    )

    if not is_image_filename(filename):
        raise ValueError("Не похоже на изображение")
    stem = Path(filename).stem
    category = category_from_image_stem(stem)
    text = ocr_image_bytes(data)
    if not category:
        raise ValueError(
            f"Укажите категорию в имени файла, например VACANCY_{stem}.jpg, "
            "или отправьте фото с подписью VACANCY",
        )
    if len(text) < 15:
        raise ValueError("OCR не нашёл достаточно текста (минимум 15 символов)")
    return [(text, category)]


def load_rows_from_zip_bytes(data: bytes) -> list[tuple[str, str]]:
    from app.ml.image_ocr import (
        category_from_image_stem,
        is_image_filename,
        ocr_image_bytes,
    )

    rows: list[tuple[str, str]] = []
    manifest_ext = (".xlsx", ".xlsm", ".csv", ".tsv", ".txt", ".json", ".jsonl", ".ndjson")
    with zipfile.ZipFile(BytesIO(data)) as archive:
        names = [name for name in archive.namelist() if not name.endswith("/")]
        manifest_files = [
            name for name in names
            if Path(name).name.lower().endswith(manifest_ext)
            and "/__macosx" not in name.lower()
        ]
        manifest_files.sort(key=lambda n: (0 if "sample" in n.lower() or "train" in n.lower() else 1, n))

        for name in manifest_files:
            try:
                rows.extend(parse_training_file(Path(name).name, archive.read(name)))
            except Exception:
                continue

        if rows:
            return _dedupe_rows(rows)

        for name in names:
            if "/__macosx" in name.lower():
                continue
            if not is_image_filename(name):
                continue
            stem = Path(name).stem
            category = category_from_image_stem(stem)
            if not category:
                continue
            try:
                text = ocr_image_bytes(archive.read(name))
            except Exception:
                continue
            if len(text) >= 15:
                rows.append((text, category))

    return _dedupe_rows(rows)


def _dedupe_rows(rows: list[tuple[str, str]]) -> list[tuple[str, str]]:
    seen: set[tuple[str, str]] = set()
    out: list[tuple[str, str]] = []
    for text, category in rows:
        key = (text.strip(), category)
        if key in seen:
            continue
        seen.add(key)
        out.append(key)
    return out


def parse_training_file(filename: str, data: bytes) -> list[tuple[str, str]]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        return _dedupe_rows(load_rows_from_excel_bytes(data))
    if name.endswith(".xls"):
        raise ValueError("Старый .xls не поддерживается — сохраните как .xlsx")
    if name.endswith(".tsv"):
        return _dedupe_rows(load_rows_from_csv_bytes(data, delimiter="\t"))
    if name.endswith(".csv"):
        return _dedupe_rows(load_rows_from_csv_bytes(data, delimiter=","))
    if name.endswith((".json",)):
        return _dedupe_rows(load_rows_from_json_bytes(data))
    if name.endswith((".jsonl", ".ndjson")):
        return _dedupe_rows(load_rows_from_jsonl_bytes(data))
    if name.endswith(".zip"):
        return load_rows_from_zip_bytes(data)
    if name.endswith((".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff")):
        return load_rows_from_image_bytes(data, filename)
    if name.endswith(".txt"):
        for encoding in ("utf-8-sig", "utf-8", "cp1251"):
            try:
                return _dedupe_rows(load_rows_from_delimited_text(data.decode(encoding)))
            except UnicodeDecodeError:
                continue
        raise ValueError("Не удалось прочитать текстовый файл (кодировка)")

    # content sniff
    if data[:2] == b"PK":
        return load_rows_from_zip_bytes(data)
    if data[:1] in (b"{", b"["):
        return _dedupe_rows(load_rows_from_json_bytes(data))
    stripped = data.lstrip()
    if stripped[:1] == b"{":
        return _dedupe_rows(load_rows_from_jsonl_bytes(data))

    raise ValueError(
        "Формат не поддерживается. Нужен .xlsx, .csv, .tsv, .txt, .json, .jsonl, .zip или изображение",
    )


def load_rows_from_path(path: Path) -> list[tuple[str, str]]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return load_rows_from_excel(path)
    return parse_training_file(path.name, path.read_bytes())


def parse_photo_caption(caption: str) -> tuple[str | None, str | None, str | None]:
    """
    Returns (category, text_override, error).
    caption modes:
      - VACANCY
      - VACANCY\\nfull ad text
      - full ad text only
    """
    caption = (caption or "").strip()
    if not caption:
        return None, None, None
    lines = caption.split("\n", 1)
    first_cat = resolve_category_token(lines[0].strip())
    if first_cat and len(lines) > 1:
        body = lines[1].strip()
        if len(body) >= 15:
            return first_cat, body, None
        return first_cat, None, None
    if first_cat and len(caption) <= 32:
        return first_cat, None, None
    if len(caption) >= 15:
        return None, caption, None
    return None, None, "Подпись слишком короткая"
